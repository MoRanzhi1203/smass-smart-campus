"""批量导入回归测试（对应计划 9.4 导入部分）。

覆盖：
- 模板表头与解析索引一致（历史缺陷：模板 10 列、解析按 13 列读，必然错位）；
- 成功导入的账号密码为官方哈希，且可用返回的统一口令登录；
- 重复主键计入 skipped，不产生重复行；
- 任一行非法 → 整批回滚，数据库无半成品；
- 非 Excel 文件给出可读错误而非 500。
"""
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from app01 import models
from app01.services import accounts, imports
from app01.tests import helpers


def build_upload(headers, rows, filename='import.xlsx'):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return SimpleUploadedFile(
        filename,
        stream.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def student_row(student_id, name, major='信息安全技术', classes='信安2101',
                sex='男', political='群众'):
    return [
        student_id, name, 'pinyin', sex, '汉族', '500000200301010001',
        '13800001234', '123456789', '重庆市永川区', political, major, classes, '', '',
    ]


def teacher_row(teacher_id, name, positions='辅导员', college='信息工程学院',
                sex='男', political='党员'):
    return [
        teacher_id, name, 'pinyin', sex, '汉族', '500000199001010088',
        '13800005678', '12345678', '重庆市永川区', '2024-09-01',
        positions, college, political,
    ]


class ImportBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin, cls.admin_password = helpers.make_admin()
        cls.college = helpers.make_college(name='信息工程学院')
        cls.grade = helpers.make_grade(name='2021级')
        cls.major = helpers.make_major(college=cls.college, name='信息安全技术')
        cls.teacher, _ = helpers.make_teacher(college=cls.college, name='张老师')
        cls.classes = helpers.make_classes(
            major=cls.major, grade=cls.grade,
            counsellor=cls.teacher, name='信安2101',
        )

    def setUp(self):
        self.client = helpers.logged_in_client('admin', self.admin.pk, self.admin_password)


class TemplateConsistencyTests(ImportBase):
    """模板与解析器必须使用同一份表头定义。"""

    def test_student_template_matches_parser(self):
        resp = self.client.get('/admin/student/download/template/')
        self.assertEqual(resp.status_code, 200)

        workbook = load_workbook(BytesIO(resp.content))
        headers = [cell.value for cell in workbook.active[1]]
        self.assertEqual(headers, imports.STUDENT_HEADERS)

    def test_teacher_template_matches_parser(self):
        resp = self.client.get('/admin/teacher/download/template/')
        self.assertEqual(resp.status_code, 200)

        workbook = load_workbook(BytesIO(resp.content))
        headers = [cell.value for cell in workbook.active[1]]
        self.assertEqual(headers, imports.TEACHER_HEADERS)

    def test_classes_template_matches_parser(self):
        resp = self.client.get('/admin/classes/download/template/')
        workbook = load_workbook(BytesIO(resp.content))
        headers = [cell.value for cell in workbook.active[1]]
        self.assertEqual(headers, imports.CLASSES_HEADERS)


class StudentImportTests(ImportBase):
    """学生导入。"""

    def test_import_success_uses_hashed_password(self):
        upload = build_upload(
            imports.STUDENT_HEADERS,
            [student_row('2024000001', '导入学生甲')],
        )
        resp = self.client.post('/admin/student/multi/', {'excel_template': upload})
        self.assertEqual(resp.status_code, 200)

        result = resp.context['result']
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.created, 1)

        student = models.Student.objects.get(pk='2024000001')
        # 绝不能是明文（历史缺陷：password='123456' 直接 bulk_create）
        self.assertNotEqual(student.password, '123456')
        self.assertTrue(student.password.startswith('pbkdf2_'))
        # 用页面展示的一次性口令可以登录
        self.assertTrue(accounts.verify_password(student, result.shared_password))

    def test_duplicate_ids_are_skipped_not_duplicated(self):
        upload = build_upload(
            imports.STUDENT_HEADERS,
            [student_row('2024000002', '第一次导入')],
        )
        self.client.post('/admin/student/multi/', {'excel_template': upload})
        self.assertEqual(models.Student.objects.filter(pk='2024000002').count(), 1)

        upload2 = build_upload(
            imports.STUDENT_HEADERS,
            [student_row('2024000002', '第二次导入')],
        )
        resp = self.client.post('/admin/student/multi/', {'excel_template': upload2})
        result = resp.context['result']

        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.created, 0)
        self.assertEqual(result.skipped, 1)
        self.assertEqual(models.Student.objects.filter(pk='2024000002').count(), 1)
        # 不覆盖原有姓名
        self.assertEqual(models.Student.objects.get(pk='2024000002').name, '第一次导入')

    def test_invalid_reference_rolls_back_whole_batch(self):
        """一行合法 + 一行专业不存在 → 整批不写入。"""
        before = models.Student.objects.count()

        upload = build_upload(imports.STUDENT_HEADERS, [
            student_row('2024000003', '合法学生'),
            student_row('2024000004', '非法学生', major='不存在的专业'),
        ])
        resp = self.client.post('/admin/student/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertFalse(result.ok)
        self.assertTrue(any('不存在的专业' in e for e in result.errors), result.errors)
        self.assertEqual(models.Student.objects.count(), before)
        self.assertFalse(models.Student.objects.filter(pk='2024000003').exists())

    def test_invalid_sex_value_is_reported(self):
        upload = build_upload(
            imports.STUDENT_HEADERS,
            [student_row('2024000005', '性别错误', sex='未知')],
        )
        resp = self.client.post('/admin/student/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertFalse(result.ok)
        self.assertTrue(any('性别' in e for e in result.errors), result.errors)
        self.assertFalse(models.Student.objects.filter(pk='2024000005').exists())

    def test_non_excel_file_gives_readable_error(self):
        bad = SimpleUploadedFile('bad.txt', b'not an excel file', content_type='text/plain')
        resp = self.client.post('/admin/student/multi/', {'excel_template': bad})
        self.assertEqual(resp.status_code, 200)

        result = resp.context['result']
        self.assertFalse(result.ok)
        self.assertTrue(any('Excel' in e for e in result.errors), result.errors)

    def test_missing_file_gives_readable_error(self):
        resp = self.client.post('/admin/student/multi/', {})
        self.assertEqual(resp.status_code, 200)
        result = resp.context['result']
        self.assertFalse(result.ok)


class TeacherImportTests(ImportBase):
    """教师导入（含 through 中间表）。"""

    def test_import_teacher_creates_positions(self):
        helpers.make_position(name='辅导员')
        helpers.make_position(name='专任教师')

        upload = build_upload(
            imports.TEACHER_HEADERS,
            [teacher_row('T910001', '导入教师', positions='辅导员,专任教师')],
        )
        resp = self.client.post('/admin/teacher/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.created, 1)

        teacher = models.Teacher.objects.get(pk='T910001')
        self.assertTrue(teacher.password.startswith('pbkdf2_'))
        self.assertTrue(accounts.verify_password(teacher, result.shared_password))

        self.assertEqual(
            models.TeacherPosition.objects.filter(teacher=teacher).count(), 2
        )

    def test_import_teacher_unknown_position_rolls_back(self):
        upload = build_upload(
            imports.TEACHER_HEADERS,
            [teacher_row('T910002', '职务错误教师', positions='不存在的职务')],
        )
        resp = self.client.post('/admin/teacher/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertFalse(result.ok)
        self.assertFalse(models.Teacher.objects.filter(pk='T910002').exists())

    def test_import_teacher_unknown_college_rolls_back(self):
        helpers.make_position(name='辅导员')
        upload = build_upload(
            imports.TEACHER_HEADERS,
            [teacher_row('T910003', '院系错误', college='不存在学院')],
        )
        resp = self.client.post('/admin/teacher/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertFalse(result.ok)
        self.assertFalse(models.Teacher.objects.filter(pk='T910003').exists())


class OrgImportTests(ImportBase):
    """学院 / 专业 / 班级 导入。"""

    def test_import_colleges_is_idempotent(self):
        upload = build_upload(imports.COLLEGE_HEADERS, [['智能制造学院']])
        resp = self.client.post('/admin/college/multi/', {'excel_template': upload})
        self.assertEqual(resp.context['result'].created, 1)

        upload2 = build_upload(imports.COLLEGE_HEADERS, [['智能制造学院']])
        resp2 = self.client.post('/admin/college/multi/', {'excel_template': upload2})
        result = resp2.context['result']

        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.created, 0)
        self.assertEqual(result.skipped, 1)
        self.assertEqual(models.College.objects.filter(name='智能制造学院').count(), 1)

    def test_import_majors_requires_existing_college(self):
        upload = build_upload(imports.MAJOR_HEADERS, [['信息工程学院', '软件技术']])
        resp = self.client.post('/admin/major/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertTrue(result.ok, result.errors)
        self.assertTrue(models.Major.objects.filter(name='软件技术').exists())

    def test_import_majors_unknown_college_rolls_back(self):
        upload = build_upload(imports.MAJOR_HEADERS, [['不存在学院', '新专业X']])
        resp = self.client.post('/admin/major/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertFalse(result.ok)
        self.assertFalse(models.Major.objects.filter(name='新专业X').exists())

    def test_import_classes_roundtrip(self):
        upload = build_upload(imports.CLASSES_HEADERS, [
            ['信息工程学院', '信息安全技术', '2021级', '信安2102', '张老师', self.teacher.pk],
        ])
        resp = self.client.post('/admin/classes/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertTrue(result.ok, result.errors)
        created = models.Classes.objects.get(name='信安2102')
        self.assertEqual(created.counsellor_id, self.teacher.pk)
        self.assertEqual(created.major_id, self.major.pk)

    def test_import_classes_unknown_teacher_rolls_back(self):
        upload = build_upload(imports.CLASSES_HEADERS, [
            ['信息工程学院', '信息安全技术', '2021级', '信安2103', '-', 'T000000'],
        ])
        resp = self.client.post('/admin/classes/multi/', {'excel_template': upload})
        result = resp.context['result']

        self.assertFalse(result.ok)
        self.assertFalse(models.Classes.objects.filter(name='信安2103').exists())
