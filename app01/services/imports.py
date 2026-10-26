"""Excel 批量导入服务。

原则（对应修复计划 R1-C.3）：

1. **先校验、后写库**：所有行先解析与校验，构造出对象列表；任一行不合法就
   整体放弃写入，并返回"第 N 行、原因"清单，绝不留半成品数据。
2. **单事务**：写库阶段包在一个 ``transaction.atomic()`` 内，中途异常整体回滚。
3. **密码安全**：批量导入不接受明文默认密码；每次导入生成一个随机口令，
   所有导入账号使用同一口令的官方哈希，口令只返回给调用方展示一次。
4. **不吞异常**：文件无法解析、缺少工作表等情况转成结构化错误，不抛 500。

表头与解析索引必须一一对应（历史版本 ``student_download_template`` 只写了 10 列，
而 ``student_multi`` 按 13 列解析，导致导入错位——本次一并校准）。
"""
import logging
from dataclasses import dataclass, field

from django.contrib.auth.hashers import make_password
from django.db import transaction
from openpyxl import load_workbook

from app01 import models
from app01.services.accounts import generate_password

logger = logging.getLogger('smass.audit')


@dataclass
class ImportResult:
    """导入结果。``errors`` 非空时表示**没有任何数据被写入**。"""

    created: int = 0
    skipped: int = 0
    errors: list = field(default_factory=list)
    shared_password: str = ''

    @property
    def ok(self):
        return not self.errors


class SheetError(ValueError):
    """Excel 文件本身不可用（非 xlsx、无工作表、列数不足等）。"""


#: 学生导入模板表头（顺序即解析索引）
STUDENT_HEADERS = [
    '学号', '姓名', '拼音', '性别', '民族', '身份证号', '手机号', 'QQ号',
    '详细地址', '政治面貌', '专业', '班级', '楼栋', '寝室号',
]

#: 教师导入模板表头
TEACHER_HEADERS = [
    '教师编号', '姓名', '拼音', '性别', '民族', '身份证号', '手机号', 'QQ号',
    '详细地址', '入职时间', '教师职务', '院系', '政治面貌',
]

#: 班级导入模板表头
CLASSES_HEADERS = ['学院', '专业', '年级', '班级', '辅导员', '教师编号']

#: 专业导入模板表头
MAJOR_HEADERS = ['学院', '专业']

#: 学院导入模板表头
COLLEGE_HEADERS = ['学院']


def _load_active_sheet(file_object):
    """读取上传文件的活动工作表。"""
    if not file_object:
        raise SheetError('未接收到文件')
    try:
        workbook = load_workbook(file_object, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 文件格式错误需转成可读提示
        raise SheetError('无法解析 Excel 文件，请上传 .xlsx 格式') from exc
    sheet = workbook.active
    if sheet is None:
        raise SheetError('工作簿中没有可用的工作表')
    return sheet


def _clean(value):
    """单元格值归一化：去首尾空白、None → ''。"""
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    return value


def _text(value):
    """转字符串并去空白；数字型单元格（如学号）会去掉多余 ``.0``。"""
    value = _clean(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_sex(value):
    text = _text(value)
    if text in ('男', '0'):
        return 0, None
    if text in ('女', '1'):
        return 1, None
    if text == '':
        return None, None
    return None, '性别只能是"男"或"女"'


def _parse_political(value):
    text = _text(value)
    if text in ('', '群众', '0'):
        return 0, None
    if text in ('团员', '1'):
        return 1, None
    if text in ('党员', '2'):
        return 2, None
    return None, '政治面貌只能是"群众/团员/党员"'


def _parse_positive_int(value):
    """解析自增枚举值（系别/职务/年级等），非数字返回 None。"""
    text = _text(value)
    if not text.isdigit():
        return None
    return int(text)


def _iter_data_rows(sheet, min_columns):
    """跳过表头，逐行产出 ``(行号, 值元组)``，并校验最小列数。"""
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row is None:
            continue
        if all(_clean(cell) == '' for cell in row):
            continue
        if len(row) < min_columns:
            raise SheetError(
                '第 {} 行只有 {} 列，模板需要至少 {} 列'.format(index, len(row), min_columns)
            )
        yield index, row


# ---------------------------------------------------------------------------
# 学院
# ---------------------------------------------------------------------------
def import_colleges(file_object):
    """批量导入学院（幂等：同名跳过）。"""
    sheet = _load_active_sheet(file_object)
    result = ImportResult()
    pending = []

    existing = set(models.College.objects.values_list('name', flat=True))
    for line, row in _iter_data_rows(sheet, 1):
        name = _text(row[0])
        if not name:
            result.errors.append('第 {} 行：学院名称不能为空'.format(line))
            continue
        if name in existing:
            result.skipped += 1
            continue
        existing.add(name)
        pending.append(models.College(name=name))

    if result.errors:
        return result

    with transaction.atomic():
        models.College.objects.bulk_create(pending)
    result.created = len(pending)
    return result


# ---------------------------------------------------------------------------
# 专业
# ---------------------------------------------------------------------------
def import_majors(file_object):
    """批量导入专业。"""
    sheet = _load_active_sheet(file_object)
    result = ImportResult()
    pending = []

    colleges = {c.name: c for c in models.College.objects.all()}
    existing = set(models.Major.objects.values_list('name', flat=True))

    for line, row in _iter_data_rows(sheet, 2):
        college_name = _text(row[0])
        name = _text(row[1])

        if not name:
            result.errors.append('第 {} 行：专业名称不能为空'.format(line))
            continue
        college = colleges.get(college_name)
        if college is None:
            result.errors.append('第 {} 行：学院"{}"不存在'.format(line, college_name))
            continue
        if name in existing:
            result.skipped += 1
            continue
        existing.add(name)
        pending.append(models.Major(college=college, name=name))

    if result.errors:
        return result

    with transaction.atomic():
        models.Major.objects.bulk_create(pending)
    result.created = len(pending)
    return result


# ---------------------------------------------------------------------------
# 班级
# ---------------------------------------------------------------------------
def import_classes(file_object):
    """批量导入班级。"""
    sheet = _load_active_sheet(file_object)
    result = ImportResult()
    pending = []

    majors = {m.name: m for m in models.Major.objects.all()}
    grades = {g.name: g for g in models.Grade.objects.all()}
    teachers = {t.pk: t for t in models.Teacher.objects.all()}
    existing = set(models.Classes.objects.values_list('major_id', 'name'))

    for line, row in _iter_data_rows(sheet, 6):
        major_name = _text(row[1])
        grade_name = _text(row[2])
        name = _text(row[3])
        teacher_id = _text(row[5])

        if not name:
            result.errors.append('第 {} 行：班级名称不能为空'.format(line))
            continue

        major = majors.get(major_name)
        if major is None:
            result.errors.append('第 {} 行：专业"{}"不存在'.format(line, major_name))
            continue

        grade = grades.get(grade_name)
        if grade is None:
            result.errors.append('第 {} 行：年级"{}"不存在'.format(line, grade_name))
            continue

        counsellor = teachers.get(teacher_id)
        if counsellor is None:
            result.errors.append('第 {} 行：教师编号"{}"不存在'.format(line, teacher_id))
            continue

        if (major.pk, name) in existing:
            result.skipped += 1
            continue
        existing.add((major.pk, name))
        pending.append(
            models.Classes(name=name, grade=grade, major=major, counsellor=counsellor)
        )

    if result.errors:
        return result

    with transaction.atomic():
        models.Classes.objects.bulk_create(pending)
    result.created = len(pending)
    return result


# ---------------------------------------------------------------------------
# 学生
# ---------------------------------------------------------------------------
def import_students(file_object):
    """批量导入学生。

    所有导入账号使用**同一个随机口令**（本次导入生成，返回一次供管理员转达），
    不再使用固定 ``123456`` 或身份证后六位。
    """
    sheet = _load_active_sheet(file_object)
    result = ImportResult()
    pending = []

    majors = {m.name: m for m in models.Major.objects.all()}
    classes_cache = {}
    for cls in models.Classes.objects.select_related('major').all():
        classes_cache.setdefault((cls.major.name if cls.major else '', cls.name), cls)

    existing_ids = set(models.Student.objects.values_list('id', flat=True))
    shared_password = generate_password()
    hashed = make_password(shared_password)

    for line, row in _iter_data_rows(sheet, len(STUDENT_HEADERS)):
        student_id = _text(row[0])
        name = _text(row[1])
        pinyin = _text(row[2])
        sex, sex_error = _parse_sex(row[3])
        nation = _text(row[4])
        card_id = _text(row[5])
        phone = _text(row[6])
        qq = _text(row[7])
        address = _text(row[8])
        political, political_error = _parse_political(row[9])
        major_name = _text(row[10])
        classes_name = _text(row[11])
        building = _text(row[12])
        dormitory = _text(row[13])

        row_errors = []
        if not student_id:
            row_errors.append('学号不能为空')
        elif len(student_id) > 10:
            row_errors.append('学号长度不能超过 10 位')
        if not name:
            row_errors.append('姓名不能为空')
        if not card_id:
            row_errors.append('身份证号不能为空')
        elif len(card_id) > 18:
            row_errors.append('身份证号长度不能超过 18 位')
        if sex_error:
            row_errors.append(sex_error)
        if political_error:
            row_errors.append(political_error)

        classes = classes_cache.get((major_name, classes_name))
        if classes is None:
            row_errors.append('专业"{}"/班级"{}"不存在'.format(major_name, classes_name))

        if student_id in existing_ids:
            # 重复学号视为"已存在"，明确计入跳过而不是静默覆盖
            if row_errors:
                result.errors.extend(
                    '第 {} 行：{}'.format(line, msg) for msg in row_errors
                )
            else:
                result.skipped += 1
            continue

        if row_errors:
            result.errors.extend('第 {} 行：{}'.format(line, msg) for msg in row_errors)
            continue

        existing_ids.add(student_id)
        pending.append(
            models.Student(
                id=student_id,
                name=name,
                password=hashed,
                pinyin=pinyin or None,
                sex=sex,
                nation=nation or None,
                classes=classes,
                cardID=card_id,
                phone=phone or None,
                qq=qq or None,
                address=address or None,
                building=building or None,
                dormitory=dormitory or None,
                political=political,
            )
        )

    if result.errors:
        return result

    with transaction.atomic():
        models.Student.objects.bulk_create(pending)

    result.created = len(pending)
    result.shared_password = shared_password
    logger.info('students_imported count=%s skipped=%s', result.created, result.skipped)
    return result


# ---------------------------------------------------------------------------
# 教师
# ---------------------------------------------------------------------------
def import_teachers(file_object):
    """批量导入教师（含教师职务中间表）。"""
    sheet = _load_active_sheet(file_object)
    result = ImportResult()
    pending_teachers = []
    pending_positions = []

    colleges = {c.name: c for c in models.College.objects.all()}
    positions = {p.name: p for p in models.Position.objects.all()}
    existing_ids = set(models.Teacher.objects.values_list('id', flat=True))

    shared_password = generate_password()
    hashed = make_password(shared_password)

    for line, row in _iter_data_rows(sheet, len(TEACHER_HEADERS)):
        teacher_id = _text(row[0])
        name = _text(row[1])
        pinyin = _text(row[2])
        sex, sex_error = _parse_sex(row[3])
        nation = _text(row[4])
        card_id = _text(row[5])
        phone = _text(row[6])
        qq = _text(row[7])
        address = _text(row[8])
        entry_time = _clean(row[9])
        position_text = _text(row[10])
        college_name = _text(row[11])
        political, political_error = _parse_political(row[12])

        row_errors = []
        if not teacher_id:
            row_errors.append('教师编号不能为空')
        elif len(teacher_id) > 7:
            row_errors.append('教师编号长度不能超过 7 位')
        if not name:
            row_errors.append('姓名不能为空')
        if not card_id:
            row_errors.append('身份证号不能为空')
        if sex_error:
            row_errors.append(sex_error)
        if political_error:
            row_errors.append(political_error)

        college = colleges.get(college_name)
        if college is None:
            row_errors.append('院系"{}"不存在'.format(college_name))

        position_objects = []
        if position_text:
            for position_name in [p.strip() for p in position_text.replace('，', ',').split(',') if p.strip()]:
                position = positions.get(position_name)
                if position is None:
                    row_errors.append('教师职务"{}"不存在'.format(position_name))
                else:
                    position_objects.append(position)

        if teacher_id in existing_ids:
            if row_errors:
                result.errors.extend('第 {} 行：{}'.format(line, msg) for msg in row_errors)
            else:
                result.skipped += 1
            continue

        if row_errors:
            result.errors.extend('第 {} 行：{}'.format(line, msg) for msg in row_errors)
            continue

        existing_ids.add(teacher_id)
        pending_teachers.append(
            models.Teacher(
                id=teacher_id,
                name=name,
                password=hashed,
                pinyin=pinyin or None,
                sex=sex,
                nation=nation or None,
                cardID=card_id,
                phone=phone or None,
                qq=qq or None,
                address=address or None,
                entry_time=entry_time or None,
                college=college,
                political=political,
            )
        )
        for position in position_objects:
            pending_positions.append(
                models.TeacherPosition(teacher_id=teacher_id, position=position)
            )

    if result.errors:
        return result

    with transaction.atomic():
        models.Teacher.objects.bulk_create(pending_teachers)
        models.TeacherPosition.objects.bulk_create(pending_positions)

    result.created = len(pending_teachers)
    result.shared_password = shared_password
    logger.info('teachers_imported count=%s skipped=%s', result.created, result.skipped)
    return result
