"""学院信息后台。"""
import urllib.parse

from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app01 import models
from app01.form.admin_inter import CollegeModelForm
from app01.services import crud, imports
from app01.utils.authz import admin_required
from app01.utils.pagination import Pagination

SEARCH_FIELDS = ('name',)

STR_DICT = {'name_ch': '学院', 'name_en': 'college'}


@admin_required
def college_list(request):
    """ 学院列表 """
    keyword = request.GET.get('search', '')
    queryset = crud.apply_search(models.College.objects.all().order_by('id'), SEARCH_FIELDS, keyword)

    page_object = Pagination(request, queryset)

    context = {
        'form': CollegeModelForm(),
        'str_dict': STR_DICT,
        'endpoints': crud.build_endpoints('college'),
        'queryset': page_object.page_queryset,
        'page_string': page_object.html(),
        **crud.build_search_meta(SEARCH_FIELDS, keyword),
    }
    return render(request, 'admin/college.html', context)


@admin_required
@require_POST
def college_add(request):
    """ 添加学院（ajax 请求） """
    form = CollegeModelForm(data=request.POST)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})
    form.save()
    return JsonResponse({'status': True})


@admin_required
def college_detail(request):
    """ 根据主键获取学院信息 """
    uid = request.GET.get('uid')
    row = crud.row_detail_dict(models.College, uid)
    if row is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)
    return JsonResponse({'status': True, 'data': row})


@admin_required
@require_POST
def college_edit(request):
    """ 编辑学院（ajax 请求） """
    uid = request.GET.get('uid')
    instance = models.College.objects.filter(pk=uid).first()
    if instance is None:
        return JsonResponse({'status': False, 'tips': '数据不存在'}, status=404)

    form = CollegeModelForm(data=request.POST, instance=instance)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})
    form.save()
    return JsonResponse({'status': True})


@admin_required
@require_POST
def college_delete(request):
    """ 删除学院（ajax 请求） """
    uid = request.GET.get('uid') or request.POST.get('uid')
    return crud.delete_by_pk(models.College, uid, request)


@admin_required
def college_download_template(request):
    """ 下载【批量添加学院】模版 """
    workbook = Workbook()
    sheet = workbook.active
    for col_num, header in enumerate(imports.COLLEGE_HEADERS, 1):
        sheet['{}1'.format(get_column_letter(col_num))] = header

    filename = '学院信息批量导入模版.xlsx'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename*=UTF-8\'\'{}'.format(
        urllib.parse.quote(filename)
    )
    workbook.save(response)
    return response


@admin_required
@require_POST
def college_multi(request):
    """ 批量添加学院（Excel 文件上传，整批原子） """
    file_object = request.FILES.get('excel_template')
    try:
        result = imports.import_colleges(file_object)
    except imports.SheetError as exc:
        result = imports.ImportResult(errors=[str(exc)])

    return render(request, 'admin/import_result.html', {
        'str_dict': STR_DICT,
        'result': result,
        'list_url': '/admin/college/list/',
        'template_url': '/admin/college/download/template/',
    })
