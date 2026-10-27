"""班级信息后台。"""
import urllib.parse

from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app01 import models
from app01.form.admin_inter import ClassesModelForm
from app01.services import accounts, crud, imports
from app01.utils.authz import admin_required
from app01.utils.pagination import Pagination

SEARCH_FIELDS = (
    'name',
    'major__name',
    'major__college__name',
    'grade__name',
    'counsellor__name',
)

STR_DICT = {'name_ch': '班级', 'name_en': 'classes'}


@admin_required
def classes_list(request):
    """ 班级列表 """
    keyword = request.GET.get('search', '')
    queryset = crud.apply_search(
        models.Classes.objects.select_related(
            'major', 'major__college', 'grade', 'counsellor'
        ).order_by('id'),
        SEARCH_FIELDS,
        keyword,
    )

    page_object = Pagination(request, queryset)

    context = {
        'form': ClassesModelForm(),
        'str_dict': STR_DICT,
        'endpoints': crud.build_endpoints('classes'),
        'queryset': page_object.page_queryset,
        'page_string': page_object.html(),
        **crud.build_search_meta(SEARCH_FIELDS, keyword),
    }
    return render(request, 'admin/classes.html', context)


@admin_required
@require_POST
def classes_add(request):
    """ 添加班级（ajax 请求） """
    form = ClassesModelForm(data=request.POST)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})
    form.save()
    return JsonResponse({'status': True})


@admin_required
def classes_detail(request):
    """ 根据主键获取班级信息 """
    uid = request.GET.get('uid')
    row = crud.row_detail_dict(models.Classes, uid)
    if row is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)
    return JsonResponse({'status': True, 'data': row})


@admin_required
@require_POST
def classes_edit(request):
    """ 编辑班级（ajax 请求） """
    uid = request.GET.get('uid')
    instance = models.Classes.objects.filter(pk=uid).first()
    if instance is None:
        return JsonResponse({'status': False, 'tips': '数据不存在'}, status=404)

    form = ClassesModelForm(data=request.POST, instance=instance)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})
    form.save()
    return JsonResponse({'status': True})


@admin_required
@require_POST
def classes_delete(request):
    """ 删除班级（ajax 请求） """
    uid = request.GET.get('uid') or request.POST.get('uid')
    return crud.delete_by_pk(models.Classes, uid, request)


@admin_required
def classes_download_template(request):
    """ 下载【批量添加班级】模版 """
    workbook = Workbook()
    sheet = workbook.active
    for col_num, header in enumerate(imports.CLASSES_HEADERS, 1):
        sheet['{}1'.format(get_column_letter(col_num))] = header

    filename = '班级信息批量导入模版.xlsx'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename*=UTF-8\'\'{}'.format(
        urllib.parse.quote(filename)
    )
    workbook.save(response)
    return response


@admin_required
@require_POST
def classes_multi(request):
    """ 批量添加班级信息（Excel 文件上传，整批原子） """
    file_object = request.FILES.get('excel_template')
    try:
        result = imports.import_classes(file_object)
    except imports.SheetError as exc:
        result = imports.ImportResult(errors=[str(exc)])

    return render(request, 'admin/import_result.html', {
        'str_dict': STR_DICT,
        'result': result,
        'list_url': '/admin/classes/list/',
        'template_url': '/admin/classes/download/template/',
    })


@admin_required
@require_POST
def all_delete(request):
    """ 清空班级表 —— HTTP 入口已停用（R1.5-B.2）。

    高危批量销毁不再允许通过浏览器触发；改用受控维护命令：
        python manage.py purge_table --model Classes --confirm
    """
    accounts.log_action(
        request, 'dangerous_blocked', None, detail='endpoint=classes.all_delete'
    )
    return JsonResponse(
        {
            'status': False,
            'error': '该操作已停用（HTTP 不提供批量清表）。'
                     '请使用维护命令：python manage.py purge_table --model Classes --confirm',
        },
        status=403,
    )
