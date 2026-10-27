"""学生信息后台。"""
import urllib.parse

from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app01 import models
from app01.form.admin_inter import StudentModelForm
from app01.services import accounts, crud, imports
from app01.utils.authz import admin_required
from app01.utils.pagination import Pagination

SEARCH_FIELDS = ('name', 'id', 'classes__name', 'classes__major__name')

STR_DICT = {'name_ch': '学生', 'name_en': 'student'}


@admin_required
def student_list(request):
    """ 学生列表 """
    keyword = request.GET.get('search', '')
    queryset = crud.apply_search(
        # 模板会访问 obj.classes.major.college，必须一次 join 到位，
        # 否则每行多一次查询（查询数基准实际捕获过这个 N+1）。
        models.Student.objects.select_related(
            'classes', 'classes__major', 'classes__major__college'
        ).order_by('id'),
        SEARCH_FIELDS,
        keyword,
    )

    page_object = Pagination(request, queryset)

    context = {
        'form': StudentModelForm(),
        'str_dict': STR_DICT,
        'endpoints': crud.build_endpoints('student'),
        'queryset': page_object.page_queryset,
        'page_string': page_object.html(),
        **crud.build_search_meta(SEARCH_FIELDS, keyword),
    }
    return render(request, 'admin/student.html', context)


@admin_required
@require_POST
def student_add(request):
    """ 添加学生（ajax 请求） """
    form = StudentModelForm(data=request.POST)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})

    account, raw_password, generated = accounts.save_account_form(
        form, models.Student, request=request
    )

    result = {'status': True}
    if generated:
        result['initial_password'] = raw_password
    return JsonResponse(result)


@admin_required
def student_detail(request):
    """ 根据主键获取学生信息（不返回密码哈希） """
    uid = request.GET.get('uid')
    row = crud.row_detail_dict(models.Student, uid, exclude=('password',))
    if row is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)
    return JsonResponse({'status': True, 'data': row})


@admin_required
@require_POST
def student_edit(request):
    """ 编辑学生（ajax 请求；密码留空表示不修改） """
    uid = request.GET.get('uid')
    instance = models.Student.objects.filter(pk=uid).first()
    if instance is None:
        return JsonResponse({'status': False, 'tips': '数据不存在'}, status=404)

    form = StudentModelForm(data=request.POST, instance=instance)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})

    accounts.save_account_form(form, models.Student, instance=instance, request=request)
    return JsonResponse({'status': True})


@admin_required
@require_POST
def student_delete(request):
    """ 删除学生（ajax 请求） """
    uid = request.GET.get('uid') or request.POST.get('uid')
    return crud.delete_by_pk(models.Student, uid, request)


@admin_required
def student_obtain(request):
    """ 获取重置密码前的确认信息（仅学号与姓名） """
    uid = request.GET.get('uid')
    row = models.Student.objects.filter(pk=uid).values('id', 'name').first()
    if row is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)
    return JsonResponse({'status': True, 'data': row})


@admin_required
@require_POST
def student_resetPass(request):
    """ 重置学生密码（随机一次性密码，返回一次） """
    uid = request.GET.get('uid') or request.POST.get('uid')
    instance = models.Student.objects.filter(pk=uid).first()
    if instance is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)

    raw_password = accounts.reset_password(instance, request=request)
    return JsonResponse({'status': True, 'password': raw_password})


@admin_required
def student_download_template(request):
    """ 下载【批量添加学生账号】模版

    表头必须与 ``app01.services.imports.STUDENT_HEADERS`` 完全一致，
    历史上二者不一致导致导入列错位。
    """
    workbook = Workbook()
    sheet = workbook.active

    for col_num, header in enumerate(imports.STUDENT_HEADERS, 1):
        sheet['{}1'.format(get_column_letter(col_num))] = header

    filename = '学生信息批量导入模版.xlsx'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename*=UTF-8\'\'{}'.format(
        urllib.parse.quote(filename)
    )
    workbook.save(response)
    return response


@admin_required
@require_POST
def student_multi(request):
    """ 批量添加学生信息（Excel 文件上传，整批原子） """
    file_object = request.FILES.get('excel_template')
    try:
        result = imports.import_students(file_object)
    except imports.SheetError as exc:
        result = imports.ImportResult(errors=[str(exc)])

    return render(request, 'admin/import_result.html', {
        'str_dict': STR_DICT,
        'result': result,
        'list_url': '/admin/student/list/',
        'template_url': '/admin/student/download/template/',
    })


@admin_required
@require_POST
def all_delete(request):
    """ 清空学生表 —— HTTP 入口已停用（R1.5-B.2）。

    高危批量销毁不再允许通过浏览器触发；改用受控维护命令：
        python manage.py purge_table --model Student --confirm
    """
    accounts.log_action(
        request, 'dangerous_blocked', None, detail='endpoint=student.all_delete'
    )
    return JsonResponse(
        {
            'status': False,
            'error': '该操作已停用（HTTP 不提供批量清表）。'
                     '请使用维护命令：python manage.py purge_table --model Student --confirm',
        },
        status=403,
    )
