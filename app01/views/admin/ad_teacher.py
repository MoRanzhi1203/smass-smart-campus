"""教师信息后台。"""
import urllib.parse

from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app01 import models
from app01.form.admin_inter import TeacherModelForm
from app01.services import accounts, crud, imports
from app01.utils.authz import admin_required
from app01.utils.pagination import Pagination

SEARCH_FIELDS = ('name', 'id', 'college__name')

STR_DICT = {'name_ch': '教师', 'name_en': 'teacher'}


@admin_required
def teacher_list(request):
    """ 教师列表 """
    keyword = request.GET.get('search', '')
    queryset = crud.apply_search(
        models.Teacher.objects.select_related('college').order_by('id'),
        SEARCH_FIELDS,
        keyword,
    )

    page_object = Pagination(request, queryset)

    context = {
        'form': TeacherModelForm(),
        'str_dict': STR_DICT,
        'endpoints': crud.build_endpoints('teacher'),
        'queryset': page_object.page_queryset,
        'page_string': page_object.html(),
        **crud.build_search_meta(SEARCH_FIELDS, keyword),
    }
    return render(request, 'admin/teacher.html', context)


@admin_required
@require_POST
def teacher_add(request):
    """ 添加教师（ajax 请求） """
    form = TeacherModelForm(data=request.POST)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})

    account, raw_password, generated = accounts.save_account_form(
        form, models.Teacher, request=request
    )

    result = {'status': True}
    if generated:
        result['initial_password'] = raw_password
    return JsonResponse(result)


@admin_required
def teacher_detail(request):
    """ 根据主键获取教师信息（不返回密码哈希） """
    uid = request.GET.get('uid')
    row = crud.row_detail_dict(models.Teacher, uid, exclude=('password',))
    if row is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)
    return JsonResponse({'status': True, 'data': row})


@admin_required
@require_POST
def teacher_edit(request):
    """ 编辑教师（ajax 请求；密码留空表示不修改） """
    uid = request.GET.get('uid')
    instance = models.Teacher.objects.filter(pk=uid).first()
    if instance is None:
        return JsonResponse({'status': False, 'tips': '数据不存在'}, status=404)

    form = TeacherModelForm(data=request.POST, instance=instance)
    if not form.is_valid():
        return JsonResponse({'status': False, 'error': form.errors})

    accounts.save_account_form(form, models.Teacher, instance=instance, request=request)
    return JsonResponse({'status': True})


@admin_required
@require_POST
def teacher_delete(request):
    """ 删除教师（ajax 请求） """
    uid = request.GET.get('uid') or request.POST.get('uid')
    return crud.delete_by_pk(models.Teacher, uid, request)


@admin_required
def teacher_obtain(request):
    """ 获取重置密码前的确认信息（仅工号与姓名） """
    uid = request.GET.get('uid')
    row = models.Teacher.objects.filter(pk=uid).values('id', 'name').first()
    if row is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)
    return JsonResponse({'status': True, 'data': row})


@admin_required
@require_POST
def teacher_resetPass(request):
    """ 重置教师密码（随机一次性密码，返回一次） """
    uid = request.GET.get('uid') or request.POST.get('uid')
    instance = models.Teacher.objects.filter(pk=uid).first()
    if instance is None:
        return JsonResponse({'status': False, 'error': '数据不存在'}, status=404)

    raw_password = accounts.reset_password(instance, request=request)
    return JsonResponse({'status': True, 'password': raw_password})


@admin_required
def teacher_download_template(request):
    """ 下载【批量添加教师信息】模版 """
    workbook = Workbook()
    sheet = workbook.active
    for col_num, header in enumerate(imports.TEACHER_HEADERS, 1):
        sheet['{}1'.format(get_column_letter(col_num))] = header

    filename = '教师信息批量导入模版.xlsx'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename*=UTF-8\'\'{}'.format(
        urllib.parse.quote(filename)
    )
    workbook.save(response)
    return response


@admin_required
@require_POST
def teacher_multi(request):
    """ 批量添加教师信息（Excel 文件上传，整批原子） """
    file_object = request.FILES.get('excel_template')
    try:
        result = imports.import_teachers(file_object)
    except imports.SheetError as exc:
        result = imports.ImportResult(errors=[str(exc)])

    return render(request, 'admin/import_result.html', {
        'str_dict': STR_DICT,
        'result': result,
        'list_url': '/admin/teacher/list/',
        'template_url': '/admin/teacher/download/template/',
    })


@admin_required
@require_POST
def all_delete(request):
    """ 清空教师表（高危操作，需显式 confirm=YES 二次确认） """
    if request.POST.get('confirm') != 'YES':
        return JsonResponse(
            {'status': False, 'error': '高危操作：需要显式提交 confirm=YES'}, status=400
        )

    deleted, _ = models.Teacher.objects.all().delete()
    accounts.log_action(request, 'bulk_delete', None, detail='model=Teacher rows={}'.format(deleted))
    return JsonResponse({'status': True, 'deleted': deleted})
