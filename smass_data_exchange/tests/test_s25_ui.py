"""S25 HTTP 层测试：可见业务流程的完整链路。

对应主任务书要求的可见序列：

    模板下载 → 选择文件 → 校验/预览 → 确认提交 → 进度/结果 → 错误下载 → 导入历史

这里用 Django 测试客户端跑完整链路（比浏览器 E2E 更稳定，且能精确断言响应体）；
浏览器级烟测在 ``tests_e2e`` 中另行覆盖。
"""
from __future__ import annotations

import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase

from smass_data_exchange import build_csv_template, build_xlsx_template, get, \
    load_builtin_adapters
from smass_organization.models import OrganizationUnit, StudentRecord

load_builtin_adapters()

DATASET = 'r4_students'
PASSWORD = 'S25Ui#2024'


def make_importer(login_name='S25UI', roles=None):
    from r21_account import services
    from r21_account.models import Account, RoleCode

    account = Account.objects.create_user(
        login_name, PASSWORD, display_name=login_name, person_type='admin')
    account.must_change_password = False
    account.save(update_fields=['must_change_password'])
    services.ensure_role_groups()
    for role in (roles or [RoleCode.SYSTEM_ADMIN]):
        services.grant_role(account, role, granted_by='test')
    return Account.objects.get(pk=account.pk)


def login(client, login_name, password=PASSWORD):
    return client.post('/r4/login/', {'login_name': login_name, 'password': password})


def xlsx_with_rows(rows) -> bytes:
    """用真实模板填写，模拟用户下载→编辑→上传。"""
    from openpyxl import load_workbook

    template = build_xlsx_template(get(DATASET))
    wb = load_workbook(io.BytesIO(template))
    ws = wb['Data']
    headers = [str(c.value or '').strip().rstrip('*').strip().lower() for c in ws[1]]
    headers = [h for h in headers if h]
    ws.delete_rows(2, 1)
    for offset, row in enumerate(rows):
        for col, code in enumerate(headers, start=1):
            ws.cell(row=2 + offset, column=col, value=row.get(code))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class VisibleFlowTests(TestCase):
    def setUp(self):
        self.account = make_importer()
        self.client = Client()
        login(self.client, self.account.login_name)
        self.unit = OrganizationUnit.objects.create(
            code='UI-K', name='界面班级', unit_type='class')

    # ---------------- 模板下载 ----------------
    def test_template_xlsx_download(self):
        resp = self.client.get('/exchange/template/{}/?fmt=xlsx'.format(DATASET))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment', resp['Content-Disposition'])
        self.assertIn('.xlsx', resp['Content-Disposition'])
        self.assertEqual(resp['X-Template-Version'], get(DATASET).version)

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ['Data', 'Instructions', 'Options'])

    def test_template_csv_download_has_bom(self):
        resp = self.client.get('/exchange/template/{}/?fmt=csv'.format(DATASET))
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.startswith(b'\xef\xbb\xbf'))
        self.assertIn('.csv', resp['Content-Disposition'])

    def test_template_bad_format_rejected(self):
        resp = self.client.get('/exchange/template/{}/?fmt=ods'.format(DATASET))
        self.assertEqual(resp.status_code, 400)

    def test_template_unknown_dataset_404(self):
        resp = self.client.get('/exchange/template/no_such/?fmt=csv')
        self.assertEqual(resp.status_code, 404)

    # ---------------- 页面 ----------------
    def test_home_lists_datasets(self):
        resp = self.client.get('/exchange/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, DATASET)
        self.assertContains(resp, '下载 Excel 模板')

    def test_import_page_has_upload_form(self):
        resp = self.client.get('/exchange/import/{}/'.format(DATASET))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'file-input')
        self.assertContains(resp, 'btn-preview')
        self.assertContains(resp, 'btn-commit')
        self.assertContains(resp, 'csrfmiddlewaretoken')

    # ---------------- 预检 ----------------
    def test_preview_endpoint_returns_ticket_and_zero_writes(self):
        before = StudentRecord.objects.count()
        content = xlsx_with_rows([{'student_no': 'UI1001', 'name': '甲',
                                   'unit_code': self.unit.code}])
        upload = SimpleUploadedFile('a.xlsx', content, content_type=(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
        resp = self.client.post('/exchange/preview/{}/'.format(DATASET), {'file': upload},
                                HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data['ok'], data.get('errors'))
        self.assertTrue(data['ticket'])
        self.assertEqual(data['summary']['would_create'], 1)
        self.assertEqual(StudentRecord.objects.count(), before, '预检不得写库')

    def test_preview_without_file_returns_400(self):
        resp = self.client.post('/exchange/preview/{}/'.format(DATASET), {},
                                HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(resp.status_code, 400)

    # ---------------- 提交 ----------------
    def test_commit_endpoint_full_journey(self):
        content = xlsx_with_rows([
            {'student_no': 'UI2001', 'name': '乙', 'unit_code': self.unit.code},
            {'student_no': 'UI2002', 'name': '丙', 'unit_code': self.unit.code},
        ])
        upload = SimpleUploadedFile('a.xlsx', content)
        preview = self.client.post('/exchange/preview/{}/'.format(DATASET),
                                   {'file': upload},
                                   HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertTrue(preview['ok'], preview.get('errors'))

        upload2 = SimpleUploadedFile('a.xlsx', content)
        resp = self.client.post('/exchange/commit/{}/'.format(DATASET), {
            'file': upload2, 'ticket': preview['ticket'],
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data['committed'], data)
        self.assertEqual(data['stats']['created'], 2)
        self.assertTrue(StudentRecord.objects.filter(student_no='UI2001').exists())

    def test_commit_without_ticket_rejected(self):
        content = xlsx_with_rows([{'student_no': 'UI2003', 'name': '丁'}])
        upload = SimpleUploadedFile('a.xlsx', content)
        resp = self.client.post('/exchange/commit/{}/'.format(DATASET), {'file': upload},
                                HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(resp.status_code, 400)

    def test_commit_with_stale_ticket_returns_409(self):
        content = xlsx_with_rows([{'student_no': 'UI2004', 'name': '戊',
                                   'unit_code': self.unit.code}])
        preview = self.client.post(
            '/exchange/preview/{}/'.format(DATASET),
            {'file': SimpleUploadedFile('a.xlsx', content)},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()

        # 换一个文件内容，但用旧令牌
        other = xlsx_with_rows([{'student_no': 'UI2005', 'name': '己',
                                 'unit_code': self.unit.code}])
        resp = self.client.post(
            '/exchange/commit/{}/'.format(DATASET),
            {'file': SimpleUploadedFile('a.xlsx', other), 'ticket': preview['ticket']},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(resp.status_code, 409)
        self.assertEqual(resp.json()['code'], 'TICKET_DIGEST_MISMATCH')

    # ---------------- 错误下载 + 历史 ----------------
    def test_failed_commit_creates_job_and_error_download(self):
        content = xlsx_with_rows([{'student_no': 'UI3001', 'name': '庚',
                                   'unit_code': 'NO-SUCH'}])
        upload = SimpleUploadedFile('a.xlsx', content)
        preview = self.client.post('/exchange/preview/{}/'.format(DATASET),
                                   {'file': upload},
                                   HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertFalse(preview['ok'])

        upload2 = SimpleUploadedFile('a.xlsx', content)
        result = self.client.post('/exchange/commit/{}/'.format(DATASET), {
            'file': upload2, 'ticket': preview['ticket'],
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertFalse(result['committed'])
        self.assertIsNotNone(result.get('job_id'))

        resp = self.client.get('/exchange/errors/{}/{}/'.format(DATASET, result['job_id']))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment', resp['Content-Disposition'])

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(resp.content))
        self.assertIn('Errors', wb.sheetnames)
        codes = {row[2] for row in wb['Errors'].iter_rows(min_row=2, values_only=True)}
        self.assertIn('UNKNOWN_UNIT', codes)

    def test_history_page_and_json(self):
        html = self.client.get('/exchange/history/{}/'.format(DATASET))
        self.assertEqual(html.status_code, 200)
        self.assertContains(html, '导入历史')

        js = self.client.get('/exchange/history/{}/'.format(DATASET),
                             HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(js.status_code, 200)
        self.assertTrue(js.json()['ok'])

    # ---------------- 导出 ----------------
    def test_export_csv_neutralizes_and_scopes(self):
        StudentRecord.objects.create(student_no='UI4001', name='=cmd|calc',
                                     unit=self.unit)
        resp = self.client.get('/exchange/export/{}/?fmt=csv'.format(DATASET))
        self.assertEqual(resp.status_code, 200)
        text = resp.content.decode('utf-8-sig')
        self.assertIn("'=cmd|calc", text, '导出必须中和公式注入')


class PermissionTests(TestCase):
    def test_anonymous_redirected_to_login(self):
        client = Client()
        for path in ('/exchange/', '/exchange/import/{}/'.format(DATASET),
                     '/exchange/template/{}/'.format(DATASET)):
            with self.subTest(path=path):
                resp = client.get(path)
                self.assertEqual(resp.status_code, 302)
                self.assertIn('/r4/login/', resp['Location'])

    def test_anonymous_cannot_preview(self):
        resp = Client().post('/exchange/preview/{}/'.format(DATASET), {})
        self.assertEqual(resp.status_code, 302)

    def test_user_without_import_permission_denied(self):
        from r21_account.models import RoleCode

        student = make_importer('S25UISTU', roles=[RoleCode.STUDENT])
        client = Client()
        login(client, student.login_name)

        content = build_csv_template(get(DATASET))
        resp = client.post(
            '/exchange/preview/{}/'.format(DATASET),
            {'file': SimpleUploadedFile('a.csv', content)},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        data = resp.json()
        self.assertFalse(data['ok'])
        self.assertTrue(any(e['code'] == 'PERMISSION_DENIED' for e in data['errors']))

    def test_export_requires_export_permission(self):
        from r21_account.models import RoleCode

        student = make_importer('S25UIEXP', roles=[RoleCode.STUDENT])
        client = Client()
        login(client, student.login_name)
        resp = client.get('/exchange/export/{}/?fmt=csv'.format(DATASET))
        self.assertEqual(resp.status_code, 403)
