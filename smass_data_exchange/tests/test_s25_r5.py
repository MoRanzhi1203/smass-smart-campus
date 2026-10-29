"""S25-R5 验收：学年 / 学期 / 课程 / 教学班 / 选课 五类导入。

重点覆盖主任务书 S25-R5 的要求：
- 五类模板均可生成；
- 外键（学年→学期→课程→教学班→账号）缺失时逐行报错，不静默挂靠；
- **选课容量校验**（含文件内已占用名额）；
- 重复选课跳过、曾退课可恢复；
- 提交走事务，失败不产生半截数据。

运行::

    python manage.py test smass_data_exchange.tests.test_s25_r5 --settings=smass.settings_r5
"""
from __future__ import annotations

import io

from django.test import TestCase

from smass_data_exchange import (
    build_csv_template,
    build_template,
    build_xlsx_template,
    commit,
    get,
    load_builtin_adapters,
    preview,
)

import unittest

from django.conf import settings as _dj_settings

_APPS = {a.split('.')[0] for a in _dj_settings.INSTALLED_APPS}
if 'smass_academic' not in _APPS:
    # 本文件验证 R5 适配器，需要 smass_academic；在只装了 R4 的 settings 下跳过
    raise unittest.SkipTest('smass_academic 未安装在当前 settings，跳过 R5 适配器测试')

load_builtin_adapters()

from smass_academic.models import (  # noqa: E402
    AcademicTerm,
    AcademicYear,
    Course,
    Enrollment,
    TeachingSection,
)

PASSWORD = 'S25R5#2024'


def make_admin(login_name='R5ADM'):
    from r21_account import services
    from r21_account.models import Account, RoleCode

    account = Account.objects.create_user(
        login_name, PASSWORD, display_name=login_name, person_type='admin')
    account.must_change_password = False
    account.save(update_fields=['must_change_password'])
    services.ensure_role_groups()
    services.grant_role(account, RoleCode.SYSTEM_ADMIN, granted_by='test')
    return Account.objects.get(pk=account.pk)


def make_student(login_name):
    from r21_account.models import Account

    account = Account.objects.create_user(
        login_name, PASSWORD, display_name=login_name, person_type='student')
    account.must_change_password = False
    account.save(update_fields=['must_change_password'])
    return Account.objects.get(pk=account.pk)


def seed_term():
    year = AcademicYear.objects.create(
        code='2024-2025', start_date='2024-09-01', end_date='2025-08-31')
    term = AcademicTerm.objects.create(
        year=year, semester='1', name='2024 秋',
        start_date='2024-09-01', end_date='2025-01-15')
    course = Course.objects.create(code='CS101', title='程序设计基础', credit=3)
    return year, term, course


def upload_result(rows, dataset, account):
    """用真实模板填写并走预检 + 提交，返回 (preview, commit)。"""
    from openpyxl import load_workbook

    schema = get(dataset)
    template = build_xlsx_template(schema)
    wb = load_workbook(io.BytesIO(template))
    ws = wb['Data']
    headers = [str(c.value or '').strip().rstrip('*').strip().lower() for c in ws[1]]
    headers = [h for h in headers if h and h != '__example__']
    ws.delete_rows(2, 1)
    for offset, row in enumerate(rows):
        for col, code in enumerate(headers, start=1):
            ws.cell(row=2 + offset, column=col, value=row.get(code))
    buffer = io.BytesIO()
    wb.save(buffer)
    content = buffer.getvalue()

    pre = preview(content, 'a.xlsx', dataset, account)
    if not pre.get('ticket'):
        return pre, None
    return pre, commit(pre['ticket'], content, 'a.xlsx', dataset, account)


class DatasetRegistrationTests(TestCase):
    def test_all_five_datasets_registered(self):
        codes = {'r5_academic_years', 'r5_academic_terms', 'r5_courses',
                 'r5_sections', 'r5_enrollments'}
        for code in codes:
            with self.subTest(code=code):
                self.assertEqual(get(code).code, code)

    def test_templates_generate_for_all_datasets(self):
        for code in ('r5_academic_years', 'r5_academic_terms', 'r5_courses',
                     'r5_sections', 'r5_enrollments'):
            with self.subTest(code=code):
                schema = get(code)
                xlsx, name, mime = build_template(schema, 'xlsx')
                self.assertTrue(xlsx)
                self.assertRegex(name, r'^[A-Za-z0-9_.\-]+$')
                csv_bytes = build_csv_template(schema)
                self.assertTrue(csv_bytes.startswith(b'\xef\xbb\xbf'))

    def test_enum_choices_come_from_models(self):
        """学期取值必须与模型一致（1/2/S），不能硬编码。"""
        schema = get('r5_academic_terms')
        semester = schema.field('semester')
        self.assertEqual(set(semester.choices), {'1', '2', 'S'})


class AcademicYearImportTests(TestCase):
    def setUp(self):
        self.admin = make_admin()

    def test_create_year(self):
        pre, res = upload_result([
            {'code': '2023-2024', 'start_date': '2023-09-01', 'end_date': '2024-08-31'},
        ], 'r5_academic_years', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['created'], 1)
        self.assertTrue(AcademicYear.objects.filter(code='2023-2024').exists())

    def test_bad_date_range_rejected(self):
        pre, res = upload_result([
            {'code': '2025-2026', 'start_date': '2026-01-01', 'end_date': '2025-01-01'},
        ], 'r5_academic_years', self.admin)
        self.assertFalse(pre['ok'])
        self.assertTrue(any(e['code'] == 'BAD_RANGE' for e in pre['errors']))
        self.assertFalse(AcademicYear.objects.filter(code='2025-2026').exists())

    def test_duplicate_in_file_rejected(self):
        pre, _res = upload_result([
            {'code': '2026-2027', 'start_date': '2026-09-01', 'end_date': '2027-08-31'},
            {'code': '2026-2027', 'start_date': '2026-09-01', 'end_date': '2027-08-31'},
        ], 'r5_academic_years', self.admin)
        self.assertTrue(any(e['code'] == 'DUPLICATE_IN_FILE' for e in pre['errors']))

    def test_idempotent_second_import_skips(self):
        rows = [{'code': '2027-2028', 'start_date': '2027-09-01',
                 'end_date': '2028-08-31'}]
        _pre1, res1 = upload_result(rows, 'r5_academic_years', self.admin)
        _pre2, res2 = upload_result(rows, 'r5_academic_years', self.admin)
        self.assertEqual(res1['stats']['created'], 1)
        self.assertEqual(res2['stats']['created'], 0)
        self.assertEqual(res2['stats']['skipped'], 1)


class TermImportTests(TestCase):
    def setUp(self):
        self.admin = make_admin()
        self.year = AcademicYear.objects.create(
            code='2024-2025', start_date='2024-09-01', end_date='2025-08-31')

    def test_unknown_year_rejected(self):
        pre, _res = upload_result([
            {'year_code': 'NO-SUCH', 'semester': '1',
             'start_date': '2024-09-01', 'end_date': '2025-01-15'},
        ], 'r5_academic_terms', self.admin)
        self.assertTrue(any(e['code'] == 'UNKNOWN_YEAR' for e in pre['errors']))

    def test_create_term(self):
        pre, res = upload_result([
            {'year_code': '2024-2025', 'semester': '2',
             'start_date': '2025-02-20', 'end_date': '2025-07-10'},
        ], 'r5_academic_terms', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['created'], 1)
        self.assertTrue(AcademicTerm.objects.filter(
            year=self.year, semester='2').exists())

    def test_invalid_semester_enum_rejected(self):
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '9',
             'start_date': '2025-02-20', 'end_date': '2025-07-10'},
        ], 'r5_academic_terms', self.admin)
        self.assertTrue(any(e['code'] == 'INVALID_VALUE' for e in pre['errors']))


class CourseImportTests(TestCase):
    def setUp(self):
        self.admin = make_admin()

    def test_create_course(self):
        pre, res = upload_result([
            {'code': 'MA101', 'title': '高等数学', 'credit': '4.5',
             'course_type': '必修'},
        ], 'r5_courses', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['created'], 1)
        course = Course.objects.get(code='MA101')
        self.assertEqual(str(course.credit), '4.5')

    def test_bad_credit_rejected(self):
        pre, _res = upload_result([
            {'code': 'MA102', 'title': '线性代数', 'credit': '-1'},
        ], 'r5_courses', self.admin)
        self.assertTrue(any(e['code'] == 'BAD_CREDIT' for e in pre['errors']))
        self.assertFalse(Course.objects.filter(code='MA102').exists())


class SectionImportTests(TestCase):
    def setUp(self):
        self.admin = make_admin()
        self.year, self.term, self.course = seed_term()

    def test_unknown_course_rejected(self):
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '1', 'course_code': 'NOPE',
             'code': 'X-01', 'capacity': '30'},
        ], 'r5_sections', self.admin)
        self.assertTrue(any(e['code'] == 'UNKNOWN_COURSE' for e in pre['errors']))

    def test_unknown_term_rejected(self):
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '2', 'course_code': 'CS101',
             'code': 'Y-01', 'capacity': '30'},
        ], 'r5_sections', self.admin)
        self.assertTrue(any(e['code'] == 'UNKNOWN_TERM' for e in pre['errors']))

    def test_create_section(self):
        pre, res = upload_result([
            {'year_code': '2024-2025', 'semester': '1', 'course_code': 'CS101',
             'code': 'CS101-01', 'capacity': '40', 'status': 'open'},
        ], 'r5_sections', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['created'], 1)
        self.assertTrue(TeachingSection.objects.filter(
            term=self.term, code='CS101-01').exists())

    def test_negative_capacity_rejected(self):
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '1', 'course_code': 'CS101',
             'code': 'CS101-02', 'capacity': '-5'},
        ], 'r5_sections', self.admin)
        self.assertTrue(any(e['code'] == 'BAD_CAPACITY' for e in pre['errors']))


class EnrollmentImportTests(TestCase):
    def setUp(self):
        self.admin = make_admin()
        self.year, self.term, self.course = seed_term()
        self.section = TeachingSection.objects.create(
            term=self.term, course=self.course, code='CS101-01',
            capacity=2, status=TeachingSection.STATUS_OPEN)
        self.s1 = make_student('R5S001')
        self.s2 = make_student('R5S002')
        self.s3 = make_student('R5S003')

    def test_enroll_creates_records(self):
        pre, res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S001'},
        ], 'r5_enrollments', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['created'], 1)
        self.assertTrue(Enrollment.objects.filter(
            section=self.section, account=self.s1).exists())

    def test_capacity_full_rejected(self):
        """容量 2：文件内 3 人，第 3 人必须被拒。"""
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S001'},
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S002'},
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S003'},
        ], 'r5_enrollments', self.admin)
        self.assertFalse(pre['ok'])
        self.assertTrue(any(e['code'] == 'SECTION_FULL' for e in pre['errors']))

    def test_capacity_respects_existing_enrollments(self):
        Enrollment.objects.create(section=self.section, account=self.s1,
                                  status='enrolled')
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S002'},
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S003'},
        ], 'r5_enrollments', self.admin)
        # 已有 1 人 + 新 1 人 = 2 人满；第 3 人（R5S003）应被拒
        self.assertTrue(any(e['code'] == 'SECTION_FULL' for e in pre['errors']))

    def test_duplicate_enrollment_skipped(self):
        Enrollment.objects.create(section=self.section, account=self.s1,
                                  status='enrolled')
        pre, res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S001'},
        ], 'r5_enrollments', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['skipped'], 1)
        self.assertEqual(Enrollment.objects.filter(
            section=self.section, account=self.s1).count(), 1)

    def test_dropped_enrollment_is_restored(self):
        Enrollment.objects.create(section=self.section, account=self.s1,
                                  status='dropped')
        pre, res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S001'},
        ], 'r5_enrollments', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['updated'], 1)
        record = Enrollment.objects.get(section=self.section, account=self.s1)
        self.assertEqual(record.status, 'enrolled')

    def test_unknown_account_rejected(self):
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'NOBODY'},
        ], 'r5_enrollments', self.admin)
        self.assertTrue(any(e['code'] == 'UNKNOWN_ACCOUNT' for e in pre['errors']))

    def test_closed_section_rejected(self):
        self.section.status = TeachingSection.STATUS_CLOSED
        self.section.save(update_fields=['status'])
        pre, _res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S001'},
        ], 'r5_enrollments', self.admin)
        self.assertTrue(any(e['code'] == 'SECTION_CLOSED' for e in pre['errors']))

    def test_zero_capacity_means_unlimited(self):
        self.section.capacity = 0
        self.section.save(update_fields=['capacity'])
        pre, res = upload_result([
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S001'},
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S002'},
            {'year_code': '2024-2025', 'semester': '1',
             'section_code': 'CS101-01', 'login_name': 'R5S003'},
        ], 'r5_enrollments', self.admin)
        self.assertTrue(pre['ok'], pre['errors'])
        self.assertEqual(res['stats']['created'], 3)


class PermissionTests(TestCase):
    def test_student_cannot_import(self):
        from r21_account.models import RoleCode

        from r21_account import services

        account = make_student('R5STU')
        services.ensure_role_groups()
        services.grant_role(account, RoleCode.STUDENT, granted_by='test')
        account = type(account).objects.get(pk=account.pk)

        pre, _res = upload_result([
            {'code': 'X1', 'title': 'X'},
        ], 'r5_courses', account)
        self.assertFalse(pre['ok'])
        self.assertTrue(any(e['code'] == 'PERMISSION_DENIED' for e in pre['errors']))
        self.assertFalse(Course.objects.filter(code='X1').exists())
