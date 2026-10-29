"""S25 数据交换中心：真实闭环与安全边界测试。

核心验收（主任务书 S25）：
- 真实下载 .xlsx / .csv 模板 → 填写 → 上传 → 预检（零业务写入）→ 提交 → 结果可查；
- 预检令牌绑定文件摘要/schema/操作者/过期，防 TOCTOU；
- 提交幂等（同文件第二次全部 skipped）；
- 错误工作簿可下载且含错误码与指引；
- 导出含公式注入中和与字段级裁剪；
- 恶意/异常输入被明确拒绝（超大、乱码、缺列、伪 xlsx、zip 炸弹、行数超限）。

运行::

    python manage.py test smass_data_exchange --settings=smass.settings_r4
"""
from __future__ import annotations

import io
import zipfile

from django.test import TestCase

from smass_data_exchange import (
    MAX_UPLOAD_BYTES,
    ParseError,
    SchemaError,
    TicketError,
    build_csv_template,
    build_error_workbook,
    build_template,
    build_xlsx_template,
    commit,
    detect_format,
    export_csv_bytes,
    export_dataset,
    get,
    load_builtin_adapters,
    parse,
    preview,
    read_ticket,
)

# 适配器依赖 smass_organization，需在 settings 已安装该 app 的环境下显式加载
import unittest

from django.conf import settings as _dj_settings

_APPS = {a.split('.')[0] for a in _dj_settings.INSTALLED_APPS}
if 'smass_organization' not in _APPS:
    # 本文件验证 R4 适配器，需要 smass_organization；在只装了 R5 的 settings 下跳过
    raise unittest.SkipTest('smass_organization 未安装在当前 settings，跳过 R4 适配器测试')

load_builtin_adapters()

from smass_organization.models import (  # noqa: E402
    UNIT_TYPE_CLASS,
    UNIT_TYPE_COLLEGE,
    DataImportJob,
    OrganizationUnit,
    StudentRecord,
)

DATASET = 'r4_students'


# ---------------------------------------------------------------------------
# 测试夹具
# ---------------------------------------------------------------------------
def make_importer(login_name='S25ADM', roles=None):
    from r21_account import services
    from r21_account.models import Account, RoleCode

    account = Account.objects.create_user(
        login_name, 'S25Pass#2024', display_name=login_name, person_type='admin')
    account.must_change_password = False
    account.save(update_fields=['must_change_password'])
    services.ensure_role_groups()
    for role in (roles or [RoleCode.SYSTEM_ADMIN]):
        services.grant_role(account, role, granted_by='test')
    return Account.objects.get(pk=account.pk)


def make_unit(code='DEMO-K', name='演示班级', unit_type=UNIT_TYPE_CLASS, parent=None):
    return OrganizationUnit.objects.create(
        code=code, name=name, unit_type=unit_type, parent=parent)


def fill_xlsx(template_bytes: bytes, data_rows) -> bytes:
    """把数据写进模板的 Data 表（覆盖示例行），返回新的 xlsx 字节。

    这就是"用真实下载的模板填写"的模拟：不改表头、不改 sheet 名。
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb['Data']
    headers = [str(c.value or '').strip().rstrip('*').strip().lower()
               for c in ws[1]]

    # 清掉示例行（第 2 行）
    ws.delete_rows(2, 1)

    for offset, row in enumerate(data_rows):
        for col, code in enumerate(headers, start=1):
            if not code:
                continue
            ws.cell(row=2 + offset, column=col, value=row.get(code))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def fill_csv(csv_bytes: bytes, data_rows) -> bytes:
    """CSV 模板：保留表头，丢弃示例行，写入真实行。"""
    text = csv_bytes.decode('utf-8-sig')
    lines = text.split('\r\n')
    header = lines[0]
    codes = [c.strip().rstrip('*').strip().lower() for c in header.split(',')]

    out = [header]
    for row in data_rows:
        values = []
        for code in codes:
            values.append('' if row.get(code) is None else str(row.get(code)))
        out.append(','.join(values))
    return ('\r\n'.join(out) + '\r\n').encode('utf-8-sig')


# ---------------------------------------------------------------------------
# 1. 模板
# ---------------------------------------------------------------------------
class TemplateTests(TestCase):
    def setUp(self):
        self.schema = get(DATASET)

    def test_xlsx_template_has_three_sheets(self):
        data = build_xlsx_template(self.schema)
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ['Data', 'Instructions', 'Options'])

    def test_xlsx_header_matches_schema(self):
        from openpyxl import load_workbook

        data = build_xlsx_template(self.schema)
        wb = load_workbook(io.BytesIO(data))
        ws = wb['Data']
        # 表头末尾有 __example__ 标记列（用于跳过示例行），比对时排除
        headers = [str(c.value or '').strip().rstrip('*').strip().lower()
                   for c in ws[1]]
        headers = [h for h in headers if h and h != '__example__']
        self.assertEqual(headers, list(self.schema.codes))

    def test_xlsx_example_row_is_synthetic(self):
        from openpyxl import load_workbook

        data = build_xlsx_template(self.schema)
        wb = load_workbook(io.BytesIO(data))
        ws = wb['Data']
        values = [c.value for c in ws[2]][:len(self.schema.codes)]
        self.assertEqual(values[0], 'DEMO-1001')
        # 示例里不得出现真实身份证/手机号格式
        joined = ' '.join('' if v is None else str(v) for v in values)
        self.assertNotRegex(joined, r'1[3-9]\d{9}')
        self.assertNotRegex(joined, r'[1-9]\d{5}(19|20)\d{2}')

    def test_xlsx_instructions_lists_every_field(self):
        from openpyxl import load_workbook

        data = build_xlsx_template(self.schema)
        wb = load_workbook(io.BytesIO(data))
        ins = wb['Instructions']
        codes = {str(row[0]) for row in ins.iter_rows(min_row=2, values_only=True)
                 if row[0]}
        for code in self.schema.codes:
            self.assertIn(code, codes)

    def test_csv_template_has_bom_and_header(self):
        data = build_csv_template(self.schema)
        self.assertTrue(data.startswith(b'\xef\xbb\xbf'), 'CSV 必须带 UTF-8 BOM')
        text = data.decode('utf-8-sig')
        header = text.split('\r\n')[0]
        codes = [c.strip().rstrip('*').strip().lower() for c in header.split(',')]
        self.assertEqual(codes, list(self.schema.codes))

    def test_template_filenames_are_ascii(self):
        for fmt in ('xlsx', 'csv'):
            _content, filename, mime = build_template(self.schema, fmt)
            self.assertRegex(filename, r'^[A-Za-z0-9_.\-]+$')
            self.assertTrue(mime)

    def test_unsupported_template_format_rejected(self):
        with self.assertRaises(ValueError):
            build_template(self.schema, 'ods')


# ---------------------------------------------------------------------------
# 2. 解析
# ---------------------------------------------------------------------------
class ParsingTests(TestCase):
    def setUp(self):
        self.schema = get(DATASET)
        self.account = make_importer()

    def test_detect_format_by_magic(self):
        self.assertEqual(detect_format('a.xlsx', build_xlsx_template(self.schema)), 'xlsx')
        self.assertEqual(detect_format('a.csv', b'student_no\nX\n'), 'csv')

    def test_fake_xlsx_rejected(self):
        with self.assertRaises(ParseError):
            detect_format('a.xlsx', b'not a zip at all')

    def test_unsupported_extension_rejected(self):
        with self.assertRaises(ParseError) as ctx:
            detect_format('a.xls', b'\xd0\xcf\x11\xe0')
        self.assertEqual(ctx.exception.code, 'UNSUPPORTED_TYPE')

    def test_oversized_upload_rejected(self):
        """大小检查属于 parse 的职责（detect_format 只看格式）。"""
        from smass_data_exchange import parse as parse_fn

        with self.assertRaises(ParseError) as ctx:
            parse_fn(b'x' * (MAX_UPLOAD_BYTES + 1), 'a.csv', self.schema)
        self.assertEqual(ctx.exception.code, 'TOO_LARGE')

    def test_missing_required_column(self):
        with self.assertRaises(ParseError) as ctx:
            parse(b'name\r\nA\r\n', 'a.csv', self.schema)
        self.assertEqual(ctx.exception.code, 'MISSING_COLUMN')

    def test_unknown_encoding_rejected(self):
        # UTF-16 字节在三种候选编码下都会产生 \x00 或解码失败
        with self.assertRaises(ParseError):
            parse(b'\xff\xfe\x00\x00student_no', 'a.csv', self.schema)

    def test_example_row_is_skipped_via_marker(self):
        """xlsx 模板的示例行由 __example__ 标记列识别并跳过。"""
        template = build_xlsx_template(self.schema)
        rows, meta, errors = parse(template, 'a.xlsx', self.schema)
        self.assertEqual(meta['skipped_example_rows'], 1)
        self.assertEqual(len(rows), 0)
        self.assertEqual(errors, [])

    def test_csv_template_has_no_example_row(self):
        """CSV 模板只含表头，避免用户必须手工删示例行。"""
        template = build_csv_template(self.schema)
        rows, meta, errors = parse(template, 'a.csv', self.schema)
        self.assertEqual(meta['skipped_example_rows'], 0)
        self.assertEqual(len(rows), 0)

    def test_user_data_equal_to_example_is_not_dropped(self):
        """用户数据恰好等于示例值时必须被导入，不能被当成示例丢弃。"""
        example = self.schema.example_row()
        csv_bytes = fill_csv(build_csv_template(self.schema), [example])
        rows, meta, errors = parse(csv_bytes, 'a.csv', self.schema)
        self.assertEqual(meta['skipped_example_rows'], 0)
        self.assertEqual(len(rows), 1, '示例同值的数据不得被丢弃')

    def test_bad_enum_value_reported(self):
        csv_bytes = fill_csv(build_csv_template(self.schema), [
            {'student_no': 'S1', 'name': '甲', 'unit_code': '', 'status': 'bogus'},
        ])
        _rows, _meta, errors = parse(csv_bytes, 'a.csv', self.schema)
        self.assertTrue(any(e['code'] == 'INVALID_VALUE' for e in errors))

    def test_gbk_csv_is_decoded(self):
        content = 'student_no,name,unit_code\r\nS9,张三,\r\n'.encode('gb18030')
        rows, meta, errors = parse(content, 'a.csv', self.schema)
        self.assertEqual(meta['encoding'], 'gb18030')
        self.assertEqual(rows[0]['name'], '张三')

    def test_zip_bomb_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
            zf.writestr('xl/big.xml', b'0' * (70 * 1024 * 1024))
        with self.assertRaises(ParseError) as ctx:
            parse(buffer.getvalue(), 'a.xlsx', self.schema)
        self.assertIn(ctx.exception.code, ('ZIP_BOMB', 'BAD_XLSX'))


# ---------------------------------------------------------------------------
# 3. 预检（零写入）
# ---------------------------------------------------------------------------
class PreviewTests(TestCase):
    def setUp(self):
        self.schema = get(DATASET)
        self.account = make_importer()
        self.unit = make_unit()

    def test_preview_writes_no_business_data(self):
        before_records = StudentRecord.objects.count()
        before_jobs = DataImportJob.objects.count()

        content = fill_xlsx(build_xlsx_template(self.schema), [
            {'student_no': 'S1001', 'name': '甲', 'unit_code': self.unit.code},
        ])
        result = preview(content, 'a.xlsx', DATASET, self.account)

        self.assertTrue(result['ok'], result['errors'])
        self.assertEqual(result['summary']['would_create'], 1)
        self.assertEqual(StudentRecord.objects.count(), before_records,
                         '预检不得写入业务数据')
        self.assertEqual(DataImportJob.objects.count(), before_jobs,
                         '预检不得写入导入作业记录')

    def test_preview_reports_unknown_unit(self):
        content = fill_xlsx(build_xlsx_template(self.schema), [
            {'student_no': 'S1002', 'name': '乙', 'unit_code': 'NO-SUCH-UNIT'},
        ])
        result = preview(content, 'a.xlsx', DATASET, self.account)
        self.assertFalse(result['ok'])
        self.assertTrue(any(e['code'] == 'UNKNOWN_UNIT' for e in result['errors']))

    def test_preview_duplicate_in_file(self):
        content = fill_xlsx(build_xlsx_template(self.schema), [
            {'student_no': 'S1003', 'name': '丙', 'unit_code': ''},
            {'student_no': 'S1003', 'name': '丁', 'unit_code': ''},
        ])
        result = preview(content, 'a.xlsx', DATASET, self.account)
        self.assertTrue(any(e['code'] == 'DUPLICATE_IN_FILE' for e in result['errors']))

    def test_preview_marks_update_for_changed_existing(self):
        StudentRecord.objects.create(student_no='S1004', name='旧名', unit=self.unit)
        content = fill_xlsx(build_xlsx_template(self.schema), [
            {'student_no': 'S1004', 'name': '新名', 'unit_code': self.unit.code},
        ])
        result = preview(content, 'a.xlsx', DATASET, self.account)
        self.assertEqual(result['summary']['would_update'], 1)

    def test_preview_marks_skip_for_identical_existing(self):
        StudentRecord.objects.create(student_no='S1005', name='同', unit=self.unit,
                                     status='enrolled')
        content = fill_xlsx(build_xlsx_template(self.schema), [
            {'student_no': 'S1005', 'name': '同', 'unit_code': self.unit.code,
             'status': 'enrolled'},
        ])
        result = preview(content, 'a.xlsx', DATASET, self.account)
        self.assertEqual(result['summary']['would_skip'], 1)

    def test_preview_denied_without_import_permission(self):
        from r21_account.models import RoleCode

        student = make_importer('S25STU', roles=[RoleCode.STUDENT])
        content = fill_csv(build_csv_template(self.schema), [
            {'student_no': 'S1006', 'name': '戊'},
        ])
        result = preview(content, 'a.csv', DATASET, student)
        self.assertFalse(result['ok'])
        self.assertTrue(any(e['code'] == 'PERMISSION_DENIED' for e in result['errors']))


# ---------------------------------------------------------------------------
# 4. 令牌（防 TOCTOU）
# ---------------------------------------------------------------------------
class TicketTests(TestCase):
    def setUp(self):
        self.schema = get(DATASET)
        self.account = make_importer()
        self.unit = make_unit()
        self.content = fill_xlsx(build_xlsx_template(self.schema), [
            {'student_no': 'S2001', 'name': '甲', 'unit_code': self.unit.code},
        ])
        self.ticket = preview(self.content, 'a.xlsx', DATASET, self.account)['ticket']

    def test_ticket_roundtrip(self):
        data = read_ticket(self.ticket)
        self.assertEqual(data.dataset, DATASET)
        self.assertEqual(data.actor, 'account:{}'.format(self.account.pk))

    def test_tampered_ticket_rejected(self):
        with self.assertRaises(TicketError) as ctx:
            read_ticket(self.ticket + 'x')
        self.assertEqual(ctx.exception.code, 'TICKET_BAD_SIGNATURE')

    def test_digest_mismatch_rejected(self):
        other = fill_xlsx(build_xlsx_template(self.schema), [
            {'student_no': 'S2002', 'name': '乙', 'unit_code': self.unit.code},
        ])
        with self.assertRaises(TicketError) as ctx:
            commit(self.ticket, other, 'a.xlsx', DATASET, self.account)
        self.assertEqual(ctx.exception.code, 'TICKET_DIGEST_MISMATCH')

    def test_actor_mismatch_rejected(self):
        other_account = make_importer('S25ADM2')
        with self.assertRaises(TicketError) as ctx:
            commit(self.ticket, self.content, 'a.xlsx', DATASET, other_account)
        self.assertEqual(ctx.exception.code, 'TICKET_ACTOR_MISMATCH')


# ---------------------------------------------------------------------------
# 5. 提交 + 幂等
# ---------------------------------------------------------------------------
class CommitTests(TestCase):
    def setUp(self):
        self.schema = get(DATASET)
        self.account = make_importer()
        self.unit = make_unit()

    def _content(self, rows):
        return fill_xlsx(build_xlsx_template(self.schema), rows)

    def test_commit_creates_records(self):
        content = self._content([
            {'student_no': 'S3001', 'name': '甲', 'unit_code': self.unit.code},
            {'student_no': 'S3002', 'name': '乙', 'unit_code': self.unit.code},
        ])
        ticket = preview(content, 'a.xlsx', DATASET, self.account)['ticket']
        result = commit(ticket, content, 'a.xlsx', DATASET, self.account)

        self.assertTrue(result['committed'], result)
        self.assertEqual(result['stats']['created'], 2)
        self.assertTrue(StudentRecord.objects.filter(student_no='S3001').exists())
        job = DataImportJob.objects.get(pk=result['job_id'])
        self.assertTrue(job.committed_at is not None)

    def test_commit_is_idempotent(self):
        rows = [{'student_no': 'S3003', 'name': '丙', 'unit_code': self.unit.code}]
        content = self._content(rows)

        first = commit(preview(content, 'a.xlsx', DATASET, self.account)['ticket'],
                       content, 'a.xlsx', DATASET, self.account)
        self.assertEqual(first['stats']['created'], 1)

        second = commit(preview(content, 'a.xlsx', DATASET, self.account)['ticket'],
                        content, 'a.xlsx', DATASET, self.account)
        self.assertEqual(second['stats']['created'], 0, '重复提交不得再次创建')
        self.assertEqual(second['stats']['skipped'], 1)
        self.assertEqual(StudentRecord.objects.filter(student_no='S3003').count(), 1)

    def test_commit_updates_changed_record(self):
        StudentRecord.objects.create(student_no='S3004', name='旧', unit=self.unit)
        content = self._content([
            {'student_no': 'S3004', 'name': '新', 'unit_code': self.unit.code},
        ])
        result = commit(preview(content, 'a.xlsx', DATASET, self.account)['ticket'],
                        content, 'a.xlsx', DATASET, self.account)
        self.assertEqual(result['stats']['updated'], 1)
        self.assertEqual(StudentRecord.objects.get(student_no='S3004').name, '新')

    def test_commit_rejected_when_errors_present(self):
        content = self._content([
            {'student_no': 'S3005', 'name': '戊', 'unit_code': 'BAD-UNIT'},
        ])
        ticket = preview(content, 'a.xlsx', DATASET, self.account)['ticket']
        result = commit(ticket, content, 'a.xlsx', DATASET, self.account)
        self.assertFalse(result['committed'])
        self.assertFalse(StudentRecord.objects.filter(student_no='S3005').exists())

    def test_csv_roundtrip(self):
        content = fill_csv(build_csv_template(self.schema), [
            {'student_no': 'S3006', 'name': '己', 'unit_code': self.unit.code},
        ])
        ticket = preview(content, 'a.csv', DATASET, self.account)['ticket']
        result = commit(ticket, content, 'a.csv', DATASET, self.account)
        self.assertTrue(result['committed'], result)
        self.assertTrue(StudentRecord.objects.filter(student_no='S3006').exists())


# ---------------------------------------------------------------------------
# 6. 错误工作簿 & 导出安全
# ---------------------------------------------------------------------------
class ErrorBookAndExportTests(TestCase):
    def setUp(self):
        self.schema = get(DATASET)

    def test_error_workbook_contains_codes(self):
        errors = [{'row': 3, 'column': 'unit_code', 'code': 'UNKNOWN_UNIT',
                   'message': '组织编码不存在'},
                  {'row': 4, 'column': 'name', 'code': 'EMPTY_NAME',
                   'message': '姓名不能为空'}]
        data = build_error_workbook(errors, DATASET)
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames[:2], ['Errors', 'Guide'])
        codes = {row[2] for row in wb['Errors'].iter_rows(min_row=2, values_only=True)}
        self.assertIn('UNKNOWN_UNIT', codes)
        self.assertIn('EMPTY_NAME', codes)

    def test_error_workbook_neutralizes_formula(self):
        errors = [{'row': 2, 'column': 'name', 'code': 'X',
                   'message': '=HYPERLINK("http://evil")'}]
        data = build_error_workbook(errors, DATASET)
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        value = wb['Errors'].cell(row=2, column=4).value
        self.assertTrue(str(value).startswith("'"), '错误工作簿也需中和公式')

    def test_export_neutralizes_formula_injection(self):
        rows = [{'student_no': 'S1', 'name': '=cmd|calc', 'unit_code': '@SUM(1)'}]
        content = export_csv_bytes(self.schema, rows)
        text = content.decode('utf-8-sig')
        self.assertIn("'=cmd|calc", text)
        self.assertIn("'@SUM(1)", text)

    def test_export_respects_allowed_columns(self):
        rows = [{'student_no': 'S1', 'name': '甲', 'unit_code': 'U',
                 'status': 'enrolled'}]
        content = export_csv_bytes(self.schema, rows, allowed_codes=['student_no'])
        header = content.decode('utf-8-sig').split('\r\n')[0]
        self.assertEqual(header, 'student_no')

    def test_export_dataset_returns_ascii_name(self):
        rows = [{'student_no': 'S1', 'name': '甲'}]
        for fmt in ('csv', 'xlsx'):
            _content, filename, mime = export_dataset(DATASET, rows, fmt=fmt)
            self.assertRegex(filename, r'^[A-Za-z0-9_.\-]+$')
            self.assertTrue(mime)


# ---------------------------------------------------------------------------
# 7. schema 约束
# ---------------------------------------------------------------------------
class SchemaTests(TestCase):
    def test_unknown_dataset_raises(self):
        with self.assertRaises(SchemaError):
            get('no_such_dataset')

    def test_duplicate_field_code_rejected(self):
        from smass_data_exchange.schema import DatasetSchema, FieldSpec

        with self.assertRaises(SchemaError):
            DatasetSchema(code='x', version='1', title='x', fields=(
                FieldSpec('a', 'A'), FieldSpec('a', 'B')))

    def test_enum_without_choices_rejected(self):
        from smass_data_exchange.schema import DatasetSchema, FieldSpec

        with self.assertRaises(SchemaError):
            DatasetSchema(code='x', version='1', title='x', fields=(
                FieldSpec('a', 'A', dtype='enum'),))
